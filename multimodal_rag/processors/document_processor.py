"""
Processors for text documents: plain text/Markdown, PDF and DOCX.
"""

import logging
import time
from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument

from ..base import BaseProcessor, ProcessingResult

logger = logging.getLogger(__name__)


class TextProcessor(BaseProcessor):
    supported_extensions = [".txt", ".md"]
    document_type = "text"

    def extract_content(self, file_path: Path) -> ProcessingResult:
        start = time.time()
        content = file_path.read_text(encoding="utf-8", errors="ignore")
        if not content.strip():
            return ProcessingResult(chunks=[], success=False, error_message="File is empty")
        chunks = self._create_chunks(content, {"word_count": len(content.split())})
        return ProcessingResult(chunks=chunks, processing_time=time.time() - start)


class PDFProcessor(BaseProcessor):
    supported_extensions = [".pdf"]
    document_type = "pdf"

    def extract_content(self, file_path: Path) -> ProcessingResult:
        start = time.time()
        parts = []
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            for page_num, page in enumerate(pdf.pages, 1):
                try:
                    # A font-relative tolerance keeps tightly kerned words apart ("SoftwareEngineer" -> "Software Engineer")
                    text = page.extract_text(x_tolerance_ratio=0.15)
                except Exception as e:
                    logger.warning(f"Failed to extract page {page_num} of {file_path.name}: {e}")
                    continue
                if text and text.strip():
                    parts.append(f"--- Page {page_num} ---\n{text}")

        if not parts:
            return ProcessingResult(
                chunks=[],
                success=False,
                error_message="PDF contains no extractable text (scanned PDFs are not supported yet)",
            )
        chunks = self._create_chunks("\n\n".join(parts), {"total_pages": total_pages})
        return ProcessingResult(chunks=chunks, processing_time=time.time() - start)


class DOCXProcessor(BaseProcessor):
    supported_extensions = [".docx"]
    document_type = "docx"

    def extract_content(self, file_path: Path) -> ProcessingResult:
        start = time.time()
        doc = DocxDocument(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        if not parts:
            return ProcessingResult(chunks=[], success=False, error_message="DOCX contains no text")
        chunks = self._create_chunks("\n".join(parts), {"title": doc.core_properties.title or ""})
        return ProcessingResult(chunks=chunks, processing_time=time.time() - start)
